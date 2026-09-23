"""投稿包路由 + 项目级新字段的往返保存。

为什么这两件事放在一个文件里：它们都是"客户端声明保存"这条链路上的坑——
`writingGenre` / `writingGoals` 属于**顶层 section 白名单**（漏一处就静默丢数据），
`publishedAt` 属于**章节内部字段**（跟着章节整体替换走，不需要白名单，但也不能被
`normalize_project` 抹掉）。两者一起测，才能证明"界面上看到的确实是存下去的那份"。
"""

from __future__ import annotations

import asyncio
import io
import zipfile

import db_gate
import pytest
from docx import Document

pytestmark = [
    pytest.mark.db,
    pytest.mark.skipif(
        not db_gate.DB_AVAILABLE,
        reason="PostgreSQL test DB unreachable (set DATABASE_URL_TEST)",
    ),
]

APP = db_gate.make_app()


def _project_payload() -> dict:
    return {
        "id": "proj-sub",
        "title": "投稿包测试",
        "genre": "轻小说 / 校园悬疑",
        "writingGenre": "novel",
        "writingGoals": {"daily": 2000, "chapter": 3000},
        "chapters": [
            {
                "id": "ch1",
                "title": "第一章",
                "prose": "他迫不急待地推开门。",
                "publishedAt": "2026-03-01T00:00:00Z",
            },
            {"id": "ch2", "title": "第二章", "prose": "第二天，雨停了。"},
        ],
    }


def _docx_text(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_submission_route_returns_docx_and_zip():
    async def _run():
        await db_gate.create_all()
        await db_gate.truncate_all()
        async with db_gate.make_client(APP) as client:
            owner = await db_gate.register_headers(client, "sub_export")
            created = await client.post(
                "/api/v1/projects", json={"title": "投稿包"}, headers=owner
            )
            pid = created.json()["id"]
            # 用整份 payload 覆盖（PUT 是整份保存）
            payload = _project_payload()
            payload["id"] = pid
            saved = await client.put(
                f"/api/v1/projects/{pid}",
                json={"data": payload, "force": True},
                headers=owner,
            )
            assert saved.status_code == 200, saved.text

            docx = await client.get(
                f"/api/v1/projects/{pid}/export/submission", headers=owner
            )
            assert docx.status_code == 200, docx.text
            assert "wordprocessingml" in docx.headers["content-type"]
            text = _docx_text(docx.content)
            assert "第一章" in text
            assert "第二章" in text
            # 投稿稿默认不带章节梗概（这里没有 synopsis，主要是确认正文在）
            assert "他迫不急待地推开门。" in text
            # 文末标注字数（默认开）
            assert "（本章 " in text

            zipped = await client.get(
                f"/api/v1/projects/{pid}/export/submission?split=1", headers=owner
            )
            assert zipped.status_code == 200, zipped.text
            assert zipped.headers["content-type"] == "application/zip"
            with zipfile.ZipFile(io.BytesIO(zipped.content)) as zf:
                names = zf.namelist()
                assert "投稿信息.txt" in names
                assert len([n for n in names if n.endswith(".docx")]) == 2

    asyncio.run(_run())


def test_export_options_are_respected():
    """缩进/梗概/字数这些开关必须真的传到排版层（否则下载下来才发现格式不对）。"""

    async def _run():
        await db_gate.create_all()
        await db_gate.truncate_all()
        async with db_gate.make_client(APP) as client:
            owner = await db_gate.register_headers(client, "sub_opts")
            created = await client.post(
                "/api/v1/projects", json={"title": "选项"}, headers=owner
            )
            pid = created.json()["id"]
            payload = _project_payload()
            payload["id"] = pid
            payload["chapters"][0]["synopsis"] = "写给自己看的梗概"
            await client.put(
                f"/api/v1/projects/{pid}",
                json={"data": payload, "force": True},
                headers=owner,
            )

            with_synopsis = await client.get(
                f"/api/v1/projects/{pid}/export/submission?synopsis=1&counts=0",
                headers=owner,
            )
            text = _docx_text(with_synopsis.content)
            assert "写给自己看的梗概" in text
            assert "（本章 " not in text  # counts=0 关掉了字数标注

            no_indent = await client.get(
                f"/api/v1/projects/{pid}/export/submission?indent=0", headers=owner
            )
            doc = Document(io.BytesIO(no_indent.content))
            body = [
                p
                for p in doc.paragraphs
                if p.text.strip() == "他迫不急待地推开门。"
            ]
            assert body, "正文段必须存在"
            assert body[0].paragraph_format.first_line_indent in (None, 0) or (
                getattr(body[0].paragraph_format.first_line_indent, "pt", 0) == 0
            )

    asyncio.run(_run())


def test_new_project_fields_round_trip_through_save():
    """writingGenre / writingGoals / publishedAt 存下去再读回来必须还在。"""

    async def _run():
        await db_gate.create_all()
        await db_gate.truncate_all()
        async with db_gate.make_client(APP) as client:
            owner = await db_gate.register_headers(client, "sub_fields")
            created = await client.post(
                "/api/v1/projects", json={"title": "字段往返"}, headers=owner
            )
            pid = created.json()["id"]
            payload = _project_payload()
            payload["id"] = pid
            await client.put(
                f"/api/v1/projects/{pid}",
                json={"data": payload, "force": True},
                headers=owner,
            )
            got = await client.get(f"/api/v1/projects/{pid}", headers=owner)
            data = got.json()
            assert data["writingGenre"] == "novel"
            # 目标模型是声明式字段，没给的项会以 None 出现——值本身回来就够了
            goals = data["writingGoals"]
            assert goals["daily"] == 2000
            assert goals["chapter"] == 3000
            assert goals.get("volume") in (None, 0)
            chapters = {c["id"]: c for c in data["chapters"]}
            assert chapters["ch1"]["publishedAt"] == "2026-03-01T00:00:00Z"
            assert not chapters["ch2"].get("publishedAt")

    asyncio.run(_run())
