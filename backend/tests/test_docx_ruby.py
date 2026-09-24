"""Word 原生注音（`w:ruby`）：结构可验证的 XML + 可回退 + 导入侧认得它。

依据：ECMA-376（WordprocessingML）的 `CT_Ruby`——`w:ruby` 依次是
`w:rubyPr`（`rubyAlign`/`hps`/`hpsRaise`/`hpsBaseText`/`lid`，顺序固定）、`w:rt`、`w:rubyBase`，
而 `w:rt` 与 `w:rubyBase` 里装的是 **run**（`w:r`）而不是裸文本。

此前判定"不做"的理由是"python-docx 不支持、也无法验证 Word 能打开这个文件"。
现在能验证的部分：XML 结构（元素/顺序/属性）与"文件还能被 python-docx 解析"；
**仍然验证不了**的是 Word 客户端的观感——所以字号是参数、并且保留了回退开关。
这两件事都写在这里，免得以后有人以为这一条被"彻底验证过"。

另一件必须钉住的事：原生注音的基准词**只在 `w:rubyBase` 里**，
所以 `paragraph.text`（只拼直接 run）会漏字——本项目的 docx 导入必须认得它，
否则"导出投稿稿 → 再导回工作台"会静默丢字。
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document

from app.core.docx_ruby import (
    DEFAULT_BASE_HPS,
    DEFAULT_LID,
    DEFAULT_RAISE_HPS,
    DEFAULT_RUBY_HPS,
    add_text,
    ruby_xml,
)
from app.core.export_submission import SubmissionOptions, submission_docx
from app.core.file_text import extract_text_from_bytes, paragraph_text
from app.core.project import normalize_project
from app.core.ruby_render import segments

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _project():
    return normalize_project(
        {
            "id": "p-ruby-docx",
            "title": "注音测试",
            "chapters": [
                {
                    "id": "c1",
                    "title": "第一章",
                    "prose": "他在校門前站住了：｜鐘《かね》が鳴った。\n{笑顔|えがお}を見せた。",
                }
            ],
        }
    )


def _document_xml(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


# ---- segments：切段（原生注音需要基准词与注音分开） ---------------------------


def test_segments_splits_plain_and_ruby_in_order():
    assert segments("前{漢字|かんじ}後") == [
        ("plain", "前", ""),
        ("ruby", "漢字", "かんじ"),
        ("plain", "後", ""),
    ]
    assert segments("｜鐘《かね》と{笑顔|えがお}") == [
        ("ruby", "鐘", "かね"),
        ("plain", "と", ""),
        ("ruby", "笑顔", "えがお"),
    ]


def test_segments_leave_incomplete_marks_as_plain_text():
    """写坏了的标记该由体检报出来，渲染层不能吞（与 render_ruby 同口径）。"""
    assert segments("｜漢字《》") == [("plain", "｜漢字《》", "")]
    assert segments("{漢字|}") == [("plain", "{漢字|}", "")]
    assert segments("") == []
    assert segments("没有注音") == [("plain", "没有注音", "")]


# ---- XML 结构（CT_Ruby） ------------------------------------------------------


def test_ruby_xml_follows_the_content_model():
    xml = ruby_xml("漢字", "かんじ")
    assert xml.startswith("<w:ruby ")
    # 三个子元素按 CT_Ruby 的顺序：rubyPr → rt → rubyBase
    order = [m.start() for m in re.finditer(r"<w:(rubyPr|rt|rubyBase)[ >]", xml)]
    assert len(order) == 3
    assert xml.index("<w:rubyPr>") < xml.index("<w:rt>") < xml.index("<w:rubyBase>")
    # rubyPr 的子元素顺序也固定
    pr = xml[xml.index("<w:rubyPr>") : xml.index("</w:rubyPr>")]
    keys = ["rubyAlign", "hps", "hpsRaise", "hpsBaseText", "lid"]
    positions = [pr.index(f"<w:{k} ") for k in keys]
    assert positions == sorted(positions)
    # 基准词与注音分别落在各自的容器里，且容器里装的是 run
    assert "<w:rubyBase><w:r>" in xml and "<w:rt><w:r>" in xml
    assert "<w:t xml:space=\"preserve\">漢字</w:t>" in xml
    assert "<w:t xml:space=\"preserve\">かんじ</w:t>" in xml


def test_ruby_xml_sizes_are_parameters_with_sane_defaults():
    xml = ruby_xml("漢字", "かんじ")
    assert f'<w:hps w:val="{DEFAULT_RUBY_HPS}"/>' in xml
    assert f'<w:hpsRaise w:val="{DEFAULT_RAISE_HPS}"/>' in xml
    assert f'<w:hpsBaseText w:val="{DEFAULT_BASE_HPS}"/>' in xml
    assert f'<w:lid w:val="{DEFAULT_LID}"/>' in xml
    assert DEFAULT_RUBY_HPS < DEFAULT_BASE_HPS, "注音字号应当小于基准词"


def test_ruby_xml_escapes_special_characters():
    xml = ruby_xml("A&B", "<x>")
    assert "A&amp;B" in xml and "&lt;x&gt;" in xml
    assert "<x>" not in xml


def test_add_text_writes_plain_runs_and_ruby_elements_in_order():
    doc = Document()
    p = doc.add_paragraph()
    add_text(p, "前{漢字|かんじ}後")
    tags = [child.tag.split("}")[1] for child in p._p]  # noqa: SLF001
    assert tags == ["r", "ruby", "r"], tags


def test_add_text_falls_back_to_rp_when_disabled():
    doc = Document()
    p = doc.add_paragraph()
    add_text(p, "|漢字《かんじ》", native=False)
    assert p.text == "漢字（かんじ）"
    assert "{" not in p.text and "《" not in p.text


# ---- 接进投稿稿 --------------------------------------------------------------


def test_submission_docx_emits_native_ruby_by_default():
    data = submission_docx(_project())
    xml = _document_xml(data)
    assert f'xmlns:w="{W_NS}"' in xml
    assert "<w:ruby>" in xml
    assert "<w:rubyBase>" in xml and "<w:rt>" in xml
    assert "鐘" in xml and "かね" in xml
    # 源标记绝不该泄漏进投稿稿
    assert "｜鐘《かね》" not in xml
    assert "{笑顔|えがお}" not in xml
    # 生成的文件必须还能被 python-docx 解析（"整份坏掉"这条要能被测出来）
    Document(io.BytesIO(data))


def test_submission_docx_can_fall_back_to_rp_text():
    data = submission_docx(_project(), SubmissionOptions(native_ruby=False))
    xml = _document_xml(data)
    assert "<w:ruby>" not in xml
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "鐘（かね）" in text


def test_submission_zip_also_carries_native_ruby():
    from app.core.export_submission import submission_zip

    with zipfile.ZipFile(io.BytesIO(submission_zip(_project()))) as zf:
        name = next(n for n in zf.namelist() if n.endswith(".docx"))
        inner = zf.read(name)
    assert "<w:ruby>" in _document_xml(inner)


# ---- 导入侧：原生注音不能变成"丢字" ------------------------------------------


def test_paragraph_text_reads_ruby_base_and_reading():
    doc = Document()
    p = doc.add_paragraph()
    add_text(p, "他在｜鐘《かね》前站住了。")
    # python-docx 的 p.text 会漏掉注音部分（这正是必须自己走一遍的原因）
    assert "鐘" not in p.text
    assert paragraph_text(p) == "他在鐘（かね）前站住了。"


def test_docx_import_keeps_annotated_words():
    """导出投稿稿 → 再导回工作台：带注音的词一个都不能丢。"""
    data = submission_docx(_project())
    text, warning = extract_text_from_bytes(data, "投稿.docx")
    assert warning is None
    assert "鐘（かね）" in text
    assert "笑顔（えがお）" in text
    assert "他在校門前站住了" in text


@pytest.mark.parametrize("native", [True, False])
def test_import_round_trip_is_stable_regardless_of_export_mode(native):
    data = submission_docx(_project(), SubmissionOptions(native_ruby=native))
    text, _ = extract_text_from_bytes(data, "投稿.docx")
    assert "鐘（かね）" in text
    assert "笑顔（えがお）" in text


def test_docx_import_still_reads_plain_hyperlinks_and_tables():
    """改造段落取文本时别把超链接与表格弄丢（原来的实现覆盖了表格）。"""
    doc = Document()
    doc.add_paragraph("普通段落")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(0, 1).text = "乙"
    buf = io.BytesIO()
    doc.save(buf)
    text, _ = extract_text_from_bytes(buf.getvalue(), "t.docx")
    assert "普通段落" in text
    assert "甲 | 乙" in text
